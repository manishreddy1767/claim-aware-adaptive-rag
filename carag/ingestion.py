"""Document and webpage ingestion.

Every extracted sentence becomes an :class:`EvidenceUnit` carrying its source
name, page number (PDFs), URL (webpages), section heading and a stable
identifier, so answers can cite exactly where evidence came from.
"""

from __future__ import annotations

import io
import logging
import re
from pathlib import Path
from urllib.parse import urlparse

from .config import IngestionConfig
from .schema import EvidenceUnit, SourceDocument
from .text_utils import clean_text, normalize_key, split_sentences

logger = logging.getLogger(__name__)

SUPPORTED_EXTENSIONS = (".pdf", ".txt", ".md", ".docx")


class IngestionError(ValueError):
    """Raised when a source cannot be read or contains no usable text."""


# Navigation / reference boilerplate commonly found on webpages.
_BOILERPLATE_PATTERNS = tuple(re.compile(p) for p in (
    r"^from wikipedia", r"^jump to", r"^for (the|other) uses", r"^contents$",
    r"^references?$", r"^external links$", r"^see also$", r"^navigation menu$",
    r"^this article is about", r"^edit$", r"^citation needed$", r"^retrieved\s",
    r"^archived from", r"^isbn\b", r"^doi\s*:", r"^issn\b", r"^main article:",
    r"^this article incorporates", r"^\[\s*\d+\s*\]", r"^(skip to|toggle|menu|search|log ?in|sign ?in|subscribe)\b",
    r"^(cookie|accept all cookies|privacy policy|terms of (use|service))",
    r"^(share|tweet|print|email) (this|article)",
))


# ---------------------------------------------------------------------------
# Block -> sentence units
# ---------------------------------------------------------------------------

def _is_heading(line: str, previous: str | None) -> bool:
    """Heuristic heading detection for PDF/TXT lines."""
    stripped = line.strip()
    if not stripped or len(stripped) > 80:
        return False
    if re.search(r"[.!?,;:]$", stripped):
        return False
    words = stripped.split()
    if len(words) > 10:
        return False
    if not (stripped[0].isupper() or stripped[0].isdigit()):
        return False
    # A heading follows a finished sentence (or starts the text).
    if previous is not None and previous.strip() and not re.search(r"[.!?:]$", previous.strip()):
        return False
    capitalized = sum(1 for w in words if w[:1].isupper() or w[:1].isdigit())
    return len(words) <= 6 or capitalized / len(words) >= 0.6


def _text_to_blocks(text: str) -> list[tuple[str, str]]:
    """Group raw lines into ('heading', text) and ('para', text) blocks."""
    blocks: list[tuple[str, str]] = []
    buffer: list[str] = []
    previous: str | None = None

    def flush() -> None:
        if buffer:
            blocks.append(("para", " ".join(buffer)))
            buffer.clear()

    for raw in text.splitlines():
        line = raw.strip()
        if not line:
            flush()
            previous = None
            continue
        if _is_heading(line, previous):
            flush()
            blocks.append(("heading", line))
            previous = line + "."  # a heading acts as a sentence boundary
            continue
        if buffer and buffer[-1].endswith("-") and line[:1].islower():
            buffer[-1] = buffer[-1][:-1] + line   # re-join hyphenated words
        else:
            buffer.append(line)
        previous = line
    flush()
    return blocks


def _split_long(sentence: str, max_chars: int) -> list[str]:
    if len(sentence) <= max_chars:
        return [sentence]
    parts = [p.strip() for p in re.split(r"(?<=;)\s+", sentence) if p.strip()]
    out: list[str] = []
    for part in parts:
        while len(part) > max_chars:
            cut = part.rfind(" ", 0, max_chars)
            cut = cut if cut > 0 else max_chars
            out.append(part[:cut].strip())
            part = part[cut:].strip()
        if part:
            out.append(part)
    return out


class _UnitBuilder:
    """Accumulates evidence units for one source, deduplicating as it goes."""

    def __init__(self, source: str, source_type: str, config: IngestionConfig,
                 url: str | None = None, web: bool = False):
        self.source = source
        self.source_type = source_type
        self.config = config
        self.url = url
        self.web = web
        self.units: list[EvidenceUnit] = []
        self.dropped = 0
        self._seen: set[str] = set()
        self._short_name = Path(source).name if source_type != "url" else (urlparse(source).netloc or source)

    def add_text(self, text: str, page: int | None = None, section: str | None = None) -> None:
        for sentence in split_sentences(text):
            for piece in _split_long(sentence, self.config.max_sentence_chars):
                self._add(piece, page, section)

    def _add(self, sentence: str, page: int | None, section: str | None) -> None:
        sentence = clean_text(sentence)
        key = normalize_key(sentence)
        words = re.findall(r"[A-Za-z0-9]+", sentence)
        if len(words) < self.config.min_sentence_words:
            self.dropped += 1
            return
        if self.web and any(p.search(key) for p in _BOILERPLATE_PATTERNS):
            self.dropped += 1
            return
        if key in self._seen:          # exact duplicate within this source
            self.dropped += 1
            return
        self._seen.add(key)
        position = len(self.units)
        page_part = f"p{page}." if page is not None else ""
        self.units.append(EvidenceUnit(
            evidence_id=f"{self._short_name}#{page_part}s{position + 1}",
            text=sentence,
            source=self.source,
            source_type=self.source_type,
            position=position,
            page=page,
            url=self.url,
            section=section,
        ))


def _finish(builder: _UnitBuilder, pages: int | None = None, title: str | None = None,
            warnings: list[str] | None = None) -> SourceDocument:
    if not builder.units:
        raise IngestionError(
            f"No extractable text found in '{builder.source}'. "
            "The file may be empty, scanned/image-only, or protected."
        )
    logger.info("Ingested %s: %d evidence units (%d dropped)",
                builder.source, len(builder.units), builder.dropped)
    return SourceDocument(
        source=builder.source,
        source_type=builder.source_type,
        units=builder.units,
        pages=pages,
        title=title,
        warnings=warnings or [],
        dropped_units=builder.dropped,
    )


# ---------------------------------------------------------------------------
# File loaders
# ---------------------------------------------------------------------------

def load_pdf(data: bytes, name: str, config: IngestionConfig) -> SourceDocument:
    from pypdf import PdfReader
    from pypdf.errors import PdfReadError

    try:
        reader = PdfReader(io.BytesIO(data))
        if reader.is_encrypted:
            try:
                reader.decrypt("")
            except Exception as exc:  # pragma: no cover - depends on file
                raise IngestionError(f"'{name}' is password protected.") from exc
        pages = list(reader.pages)
    except PdfReadError as exc:
        raise IngestionError(f"'{name}' is not a readable PDF: {exc}") from exc

    builder = _UnitBuilder(name, "pdf", config)
    warnings: list[str] = []
    empty_pages = []
    section: str | None = None
    for page_number, page in enumerate(pages, start=1):
        try:
            text = page.extract_text() or ""
        except Exception as exc:  # pypdf can fail on individual malformed pages
            warnings.append(f"Page {page_number}: text extraction failed ({exc}).")
            continue
        if not text.strip():
            empty_pages.append(page_number)
            continue
        for kind, block in _text_to_blocks(text):
            if kind == "heading":
                section = clean_text(block)
            else:
                builder.add_text(block, page=page_number, section=section)
    if empty_pages:
        warnings.append(
            f"No text on page(s) {', '.join(map(str, empty_pages))} "
            "(possibly scanned images; OCR is not supported)."
        )
    return _finish(builder, pages=len(pages), warnings=warnings)


def load_txt(data: bytes, name: str, config: IngestionConfig) -> SourceDocument:
    try:
        text = data.decode("utf-8")
    except UnicodeDecodeError:
        text = data.decode("latin-1")
    builder = _UnitBuilder(name, "txt", config)
    section = None
    for kind, block in _text_to_blocks(text):
        if kind == "heading":
            section = clean_text(block.lstrip("# "))
        else:
            builder.add_text(block, section=section)
    return _finish(builder)


def load_docx(data: bytes, name: str, config: IngestionConfig) -> SourceDocument:
    try:
        import docx
    except ImportError as exc:  # pragma: no cover
        raise IngestionError("DOCX support requires the 'python-docx' package.") from exc
    try:
        document = docx.Document(io.BytesIO(data))
    except Exception as exc:
        raise IngestionError(f"'{name}' is not a readable DOCX file: {exc}") from exc

    builder = _UnitBuilder(name, "docx", config)
    section = None
    for paragraph in document.paragraphs:
        text = paragraph.text.strip()
        if not text:
            continue
        style = (paragraph.style.name or "").lower() if paragraph.style is not None else ""
        if style.startswith("heading") or style == "title":
            section = clean_text(text)
            continue
        builder.add_text(text, section=section)
    for table in document.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells if c.text.strip()]
            if cells:
                builder.add_text("; ".join(dict.fromkeys(cells)) + ".", section=section)
    return _finish(builder)


_LOADERS = {".pdf": load_pdf, ".txt": load_txt, ".md": load_txt, ".docx": load_docx}


def load_bytes(data: bytes, name: str, config: IngestionConfig | None = None) -> SourceDocument:
    """Ingest an in-memory file (used by the web UI for uploads)."""
    config = config or IngestionConfig()
    extension = Path(name).suffix.lower()
    loader = _LOADERS.get(extension)
    if loader is None:
        raise IngestionError(
            f"Unsupported file type '{extension or name}'. "
            f"Supported: {', '.join(SUPPORTED_EXTENSIONS)} and http(s) URLs."
        )
    if not data:
        raise IngestionError(f"'{name}' is empty.")
    return loader(data, name, config)


def load_file(path: str | Path, config: IngestionConfig | None = None) -> SourceDocument:
    path = Path(path)
    if not path.exists():
        raise IngestionError(f"File not found: {path}")
    if not path.is_file():
        raise IngestionError(f"Not a file: {path}")
    return load_bytes(path.read_bytes(), path.name, config)


# ---------------------------------------------------------------------------
# Webpages
# ---------------------------------------------------------------------------

def validate_url(url: str) -> str:
    url = url.strip()
    parsed = urlparse(url)
    if (parsed.scheme not in ("http", "https") or not parsed.netloc
            or ("." not in parsed.netloc and parsed.hostname != "localhost")):
        raise IngestionError(f"Invalid URL '{url}'. Please provide a full http(s):// address.")
    return url


def html_to_document(html: str, url: str, config: IngestionConfig | None = None) -> SourceDocument:
    """Extract article text from HTML, skipping navigation and reference lists."""
    from bs4 import BeautifulSoup

    config = config or IngestionConfig()
    soup = BeautifulSoup(html, "html.parser")
    title = soup.title.get_text(strip=True) if soup.title else None

    for tag in soup(["script", "style", "nav", "footer", "header", "aside", "form",
                     "button", "noscript", "svg", "iframe", "template"]):
        tag.decompose()
    for selector in (".reflist", ".references", ".mw-references-wrap", ".navbox",
                     ".metadata", ".catlinks", ".printfooter", ".mw-editsection",
                     ".infobox", ".sidebar", ".toc", "#toc", ".mw-jump-link",
                     "[role=navigation]", "[aria-hidden=true]", ".cookie", ".advert"):
        for tag in soup.select(selector):
            tag.decompose()

    main = (soup.find("article") or soup.find("main") or soup.find(id="mw-content-text")
            or soup.body or soup)

    builder = _UnitBuilder(url, "url", config, url=url, web=True)
    block_tags = ["h1", "h2", "h3", "h4", "h5", "h6", "p", "li", "blockquote", "pre", "dd", "figcaption"]
    section = None
    found_blocks = False
    for element in main.find_all(block_tags):
        # Skip containers whose text is emitted by a nested block element.
        if element.name in ("li", "blockquote", "dd") and element.find(["p", "li"]):
            continue
        text = clean_text(element.get_text(" ", strip=True))
        if not text:
            continue
        found_blocks = True
        if element.name.startswith("h"):
            section = text
            continue
        builder.add_text(text, section=section)

    if not found_blocks:  # pages without semantic markup
        builder.add_text(main.get_text(" ", strip=True))
    return _finish(builder, title=title)


def load_url(url: str, config: IngestionConfig | None = None) -> SourceDocument:
    import requests

    config = config or IngestionConfig()
    url = validate_url(url)
    try:
        response = requests.get(
            url,
            timeout=config.request_timeout,
            headers={"User-Agent": "Mozilla/5.0 (compatible; ClaimAwareRAG/1.0; research prototype)"},
            stream=True,
        )
        response.raise_for_status()
        content_type = response.headers.get("Content-Type", "").lower()
        raw = response.raw.read(config.max_download_bytes + 1, decode_content=True)
    except requests.exceptions.RequestException as exc:
        raise IngestionError(f"Could not fetch '{url}': {exc}") from exc
    if len(raw) > config.max_download_bytes:
        raise IngestionError(f"'{url}' is larger than {config.max_download_bytes // 1_000_000} MB.")

    if "pdf" in content_type or url.lower().endswith(".pdf"):
        document = load_pdf(raw, url, config)
        for unit in document.units:
            unit.url = url
            unit.source_type = "url"
        return document
    if content_type and not any(t in content_type for t in ("html", "xml", "text/plain")):
        raise IngestionError(f"Unsupported content type '{content_type}' at '{url}'.")
    encoding = response.encoding or response.apparent_encoding or "utf-8"
    html = raw.decode(encoding, errors="replace")
    if "text/plain" in content_type:
        document = load_txt(html.encode("utf-8"), url, config)
        for unit in document.units:
            unit.url = url
            unit.source_type = "url"
        return document
    return html_to_document(html, url, config)


def load_source(source: str, config: IngestionConfig | None = None) -> SourceDocument:
    """Load a file path or an http(s) URL."""
    if source.lower().startswith(("http://", "https://")):
        return load_url(source, config)
    return load_file(source, config)
